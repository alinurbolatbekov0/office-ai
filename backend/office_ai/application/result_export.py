"""Render a stored agent result to a task-appropriate, non-executable Office file."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from office_ai.infrastructure.database import LocalDatabase

MAX_EXPORT_TEXT_CHARACTERS = 1_000_000
MAX_PRESENTATION_TEXT_CHARACTERS = 80_000
MAX_PRESENTATION_SLIDES = 50
_AGENT_FORMATS = {
    "document-management": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "business-correspondence": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "excel-analytics": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "powerpoint-presentations": ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}
_AGENT_TITLES = {
    "document-management": "Документооборот",
    "business-correspondence": "Деловая переписка",
    "excel-analytics": "Excel-аналитика",
    "powerpoint-presentations": "Презентация PowerPoint",
}


@dataclass(frozen=True)
class RenderedAgentResult:
    filename: str
    media_type: str
    content: bytes


class ResultExportError(ValueError):
    """A stored result cannot be rendered safely for its agent."""


class AgentResultExporter:
    def __init__(self, database: LocalDatabase) -> None:
        self.database = database

    def render(self, history_id: str) -> RenderedAgentResult:
        result = self.database.get_agent_task_result(history_id)
        history = self.database.get_agent_task_history_entry(history_id)
        if result is None or history is None:
            raise ResultExportError("Для этой задачи нет сохранённого результата.")
        agent_id = history["agentId"]
        output = _AGENT_FORMATS.get(agent_id)
        if output is None:
            raise ResultExportError("Для этого агента формат экспорта пока не поддерживается.")
        text = result["resultText"]
        if not isinstance(text, str) or not text.strip() or len(text) > MAX_EXPORT_TEXT_CHARACTERS:
            raise ResultExportError("Результат пуст или превышает ограничение экспорта.")
        extension, media_type = output
        if extension == "docx":
            content = _render_docx(agent_id, text)
        elif extension == "xlsx":
            content = _render_xlsx(text)
        else:
            content = _render_pptx(text)
        return RenderedAgentResult(
            filename=f"office-ai-result-{result['id']}.{extension}",
            media_type=media_type,
            content=content,
        )


def _render_docx(agent_id: str, text: str) -> bytes:
    document = Document()
    document.add_heading(_AGENT_TITLES[agent_id], level=0)
    for line in text.splitlines():
        if not line.strip():
            document.add_paragraph()
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_xlsx(text: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат"
    sheet.append(["№", "Вывод агента"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:B1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5548C8")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    chunks = [
        part
        for line in text.splitlines() or [text]
        for part in _split_text(line, 30_000)
    ]
    for index, chunk in enumerate(chunks, start=1):
        sheet.append([index, chunk])
        output_cell = sheet.cell(index + 1, 2)
        output_cell.data_type = "s"  # Never turn model text such as "=1+1" into a formula.
        output_cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 100
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _render_pptx(text: str) -> bytes:
    if len(text) > MAX_PRESENTATION_TEXT_CHARACTERS:
        raise ResultExportError("Текст слишком велик для безопасного экспорта презентации.")
    chunks = _presentation_sections(text)
    if len(chunks) > MAX_PRESENTATION_SLIDES:
        raise ResultExportError("Результат требует слишком много слайдов; сократите его и повторите задачу.")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, chunk in enumerate(chunks, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        title = slide.shapes.title
        body = slide.placeholders[1].text_frame
        heading, body_lines = chunk
        heading = heading[:120] or f"Основное содержание {index}"
        title.text = heading
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(67, 59, 157)
        title.text_frame.paragraphs[0].font.size = Pt(28)
        body.clear()
        for paragraph_text in body_lines:
            if not paragraph_text.strip():
                continue
            paragraph = body.add_paragraph()
            paragraph.text = paragraph_text.strip().lstrip("-•* ")
            paragraph.level = 0
            paragraph.font.size = Pt(19)
            paragraph.alignment = PP_ALIGN.LEFT
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


_SLIDE_HEADING = re.compile(r"^(?:#+\s*)?Слайд\s*\d+\s*(?:[-–—:.]\s*|\s+)(.+)$", re.IGNORECASE)
_MARKDOWN_HEADING = re.compile(r"^#{1,3}\s+(.+)$")


def _presentation_sections(text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    current_title: str | None = None
    current_lines: list[str] = []
    saw_explicit_heading = False
    for line in text.splitlines():
        slide_match = _SLIDE_HEADING.match(line.strip())
        markdown_match = _MARKDOWN_HEADING.match(line.strip())
        title_match = slide_match or markdown_match
        if title_match:
            if current_title is not None:
                sections.append((current_title, current_lines))
            current_title = title_match.group(1).strip()
            current_lines = []
            saw_explicit_heading = True
        elif current_title is not None:
            if line.strip():
                current_lines.append(line.strip().lstrip("-•* "))
        elif line.strip():
            current_lines.append(line.strip())
    if current_title is not None:
        sections.append((current_title, current_lines))
    if saw_explicit_heading:
        return sections

    paragraphs = [line.strip().lstrip("-•* ") for line in text.splitlines() if line.strip()]
    if not paragraphs:
        paragraphs = [text.strip()]
    sections = []
    for index, chunk in enumerate(_split_text("\n".join(paragraphs), 1_200), start=1):
        lines = chunk.splitlines() or [chunk]
        if len(lines) == 1:
            sections.append((lines[0][:120], [lines[0]]))
        else:
            sections.append((lines[0][:120], lines[1:]))
    return sections


def _split_text(text: str, limit: int) -> list[str]:
    if len(text) <= limit:
        return [text]
    chunks: list[str] = []
    current = ""
    for line in text.splitlines():
        while len(line) > limit:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(line[:limit])
            line = line[limit:]
        candidate = f"{current}\n{line}" if current else line
        if len(candidate) > limit:
            chunks.append(current)
            current = line
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks or [""]
