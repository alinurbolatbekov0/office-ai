from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from office_ai.application.result_export import AgentResultExporter, ResultExportError
from office_ai.infrastructure.database import LocalDatabase


class AgentResultExportTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_directory.cleanup)
        self.database = LocalDatabase(Path(temporary_directory.name))
        self.database.initialize()

    def _save_result(self, agent_id: str, text: str) -> str:
        history = self.database.record_agent_preview(agent_id, "synthetic-task", 1)
        self.database.save_agent_task_result(history["id"], "openai", "gpt-4.1-mini", text)
        return history["id"]

    def test_document_agents_export_text_as_docx_without_active_content(self) -> None:
        history_id = self._save_result("business-correspondence", "# Черновик\n- Синтетический ответ")
        rendered = AgentResultExporter(self.database).render(history_id)

        document = Document(BytesIO(rendered.content))
        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertTrue(rendered.filename.endswith(".docx"))
        self.assertIn("Синтетический ответ", content)
        self.assertIn("Деловая переписка", content)
        self.assertEqual(len(document.part.package.parts), len(set(part.partname for part in document.part.package.parts)))

    def test_excel_export_keeps_generated_text_as_text_not_formula(self) -> None:
        history_id = self._save_result("excel-analytics", "=1+1\nВывод: синтетические данные")
        rendered = AgentResultExporter(self.database).render(history_id)

        workbook = load_workbook(BytesIO(rendered.content), data_only=False)
        sheet = workbook["Результат"]
        self.assertTrue(rendered.filename.endswith(".xlsx"))
        self.assertEqual(sheet["B2"].value, "=1+1")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertIn("синтетические данные", sheet["B3"].value)

    def test_presentation_export_contains_only_text_shapes(self) -> None:
        history_id = self._save_result(
            "powerpoint-presentations",
            "Слайд 1 — План\n- Первый тезис\n- Второй тезис\nСлайд 2 — Итоги\n- Синтетический вывод",
        )
        rendered = AgentResultExporter(self.database).render(history_id)

        presentation = Presentation(BytesIO(rendered.content))
        slide_text = [shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertTrue(rendered.filename.endswith(".pptx"))
        self.assertTrue(any("План" in text for text in slide_text))
        self.assertTrue(any("Первый тезис" in text for text in slide_text))
        self.assertEqual(len(presentation.slides), 2)

    def test_missing_result_and_oversized_presentation_are_rejected(self) -> None:
        history = self.database.record_agent_preview("powerpoint-presentations", "synthetic-task", 1)
        with self.assertRaises(ResultExportError):
            AgentResultExporter(self.database).render(history["id"])
        self.database.save_agent_task_result(history["id"], "openai", "gpt-4.1-mini", "x" * 80_001)
        with self.assertRaises(ResultExportError):
            AgentResultExporter(self.database).render(history["id"])


if __name__ == "__main__":
    unittest.main()
